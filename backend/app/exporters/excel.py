"""Exportación a Excel (.xlsx): una hoja por grupo, por maestro y el horario general."""
from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..services.context import Ctx
from ..services.validation import load_summary
from .common import Table, general_tables, group_table, teacher_table

THIN = Side(style="thin", color="5B6B80")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SIDE_FILL = PatternFill("solid", fgColor="E8EDF5")


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "", name)[:28] or "Hoja"
    candidate, i = base, 2
    while candidate in used:
        candidate = f"{base[:25]} {i}"
        i += 1
    used.add(candidate)
    return candidate


def _write(ws, t: Table, start_row: int = 1) -> int:
    ws.cell(row=start_row, column=1, value=t.title).font = Font(bold=True, size=16)
    ws.cell(row=start_row + 1, column=1, value=t.subtitle).font = Font(size=12, color="444444")
    r0 = start_row + 3
    for j, h in enumerate(t.header, start=1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF", size=12), HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, row in enumerate(t.rows, start=1):
        ws.row_dimensions[r0 + i].height = 42
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=r0 + i, column=j, value=val or None)
            c.border = BORDER
            c.font = Font(size=12, bold=j == 1)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if j == 1:
                c.fill = SIDE_FILL
    ws.column_dimensions["A"].width = 16
    for j in range(2, len(t.header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22
    return r0 + len(t.rows) + 2


def render_excel(ctx: Ctx) -> bytes:
    wb = Workbook()
    used: set[str] = set()
    ws = wb.active
    ws.title = _sheet_name("General", used)
    row = 1
    for t in general_tables(ctx):
        row = _write(ws, t, row) + 1
    for g in sorted(ctx.groups.values(), key=lambda g: (g.grade, g.name)):
        if g.active:
            _write(wb.create_sheet(_sheet_name(f"Grupo {g.name}", used)), group_table(ctx, g.id))
    for t in sorted(ctx.teachers.values(), key=lambda t: t.code):
        if t.active:
            _write(wb.create_sheet(_sheet_name(f"{t.code} {t.name}", used)), teacher_table(ctx, t.id))

    ws = wb.create_sheet(_sheet_name("Cargas", used))
    headers = ["Identificador", "Maestro", "Carga declarada", "Horas asignadas", "Diferencia", "Horas en horario", "Estado"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
    for i, sm in enumerate(load_summary(ctx), start=2):
        for j, v in enumerate([sm.teacher_code, sm.teacher_name, sm.declared, sm.assigned, sm.difference,
                               sm.scheduled, sm.message], start=1):
            ws.cell(row=i, column=j, value=v)
    for j, w in enumerate([14, 32, 16, 16, 12, 16, 60], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
