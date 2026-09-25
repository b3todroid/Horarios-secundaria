"""Respaldos JSON y exportaciones PDF/Excel."""
import io
import json

from openpyxl import load_workbook


def _demo(client):
    pid = client.post("/api/demo/solvable").json()["id"]
    client.post(f"/api/projects/{pid}/generate", json={})
    return pid


def test_backup_roundtrip_new_project(client):
    pid = _demo(client)
    r = client.get(f"/api/projects/{pid}/backup")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    data = r.json()
    assert data["format"] == "horarios-secundaria" and len(data["teachers"]) == 5
    assert len(data["schedule"]) == 110
    files = {"file": ("respaldo.json", r.content, "application/json")}
    summary = client.post("/api/backup/inspect", files=files).json()
    assert summary["teachers"] == 5 and summary["groups"] == 5 and summary["subjects"] == 6
    restored = client.post("/api/backup/restore?mode=new", files=files).json()
    assert restored["id"] != pid and restored["is_demo"] is False
    again = client.get(f"/api/projects/{restored['id']}/backup").json()
    for key in ("teachers", "groups", "subjects", "assignments", "schedule", "availability"):
        assert len(again[key]) == len(data[key]), key
    assert again["periods"] == data["periods"]
    names = sorted(t["name"] for t in again["teachers"])
    assert names == sorted(t["name"] for t in data["teachers"])


def test_restore_replace_requires_confirmation(client, api):
    pid = _demo(client)
    backup = client.get(f"/api/projects/{pid}/backup").content
    files = {"file": ("r.json", backup, "application/json")}
    r = client.post(f"/api/backup/restore?mode=replace&project_id={api.pid}", files=files)
    assert r.status_code == 422 and "Confirma" in r.json()["detail"]
    r = client.post(f"/api/backup/restore?mode=replace&project_id={api.pid}&confirm=true", files=files)
    assert r.status_code == 200 and r.json()["id"] == api.pid
    assert len(client.get(api.url("/teachers")).json()) == 5


def test_invalid_backup_rejected(client):
    files = {"file": ("r.json", b"{not json", "application/json")}
    assert client.post("/api/backup/inspect", files=files).status_code == 422
    files = {"file": ("r.json", json.dumps({"foo": 1}).encode(), "application/json")}
    r = client.post("/api/backup/inspect", files=files)
    assert r.status_code == 422 and "respaldo" in r.json()["detail"]


def test_demo_data_is_isolated(client, api):
    api.teacher("M01", "Real")
    demo = client.post("/api/demo/solvable").json()
    assert demo["is_demo"]
    assert len(client.get(api.url("/teachers")).json()) == 1
    # Recrear la demo no duplica proyectos
    client.post("/api/demo/solvable")
    projects = client.get("/api/projects").json()
    assert sum(1 for p in projects if p["is_demo"]) == 1


def test_pdf_exports(client):
    pid = _demo(client)
    groups = client.get(f"/api/projects/{pid}/groups").json()
    teachers = client.get(f"/api/projects/{pid}/teachers").json()
    for q in ("scope=general", f"scope=group&item_id={groups[0]['id']}", "scope=group",
              f"scope=teacher&item_id={teachers[0]['id']}&orientation=portrait"):
        r = client.get(f"/api/projects/{pid}/export/pdf?{q}")
        assert r.status_code == 200 and r.content.startswith(b"%PDF"), q
        assert b"Libre" not in r.content


def test_pdf_letter_size(client):
    pid = _demo(client)
    r = client.get(f"/api/projects/{pid}/export/pdf?scope=general&orientation=portrait")
    assert b"/MediaBox [ 0 0 612 792 ]" in r.content
    r = client.get(f"/api/projects/{pid}/export/pdf?scope=general&orientation=landscape")
    assert b"/MediaBox [ 0 0 792 612 ]" in r.content


def test_excel_export(client):
    pid = _demo(client)
    r = client.get(f"/api/projects/{pid}/export/xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "General" in wb.sheetnames and "Grupo 1A" in wb.sheetnames and "Cargas" in wb.sheetnames
    values = [c.value for row in wb["Grupo 1A"].iter_rows() for c in row if c.value]
    assert "Libre" not in values and any("Español" in str(v) for v in values)


def test_json_export(client):
    pid = _demo(client)
    r = client.get(f"/api/projects/{pid}/export/json")
    assert r.json()["format"] == "horarios-secundaria"
